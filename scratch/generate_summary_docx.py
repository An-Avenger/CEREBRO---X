from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import os

def create_summary_doc():
    doc = Document()

    # Title
    title = doc.add_heading('Cerebro X: Project Progress Summary', 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # Introduction
    doc.add_heading('1. Project Overview', level=1)
    doc.add_paragraph(
        "Cerebro X is a research project aimed at creating an 'Explainable Multimodal AI-Based Digital Brain Twin for "
        "Longitudinal Prediction of Alzheimer’s Disease Progression.' Our goal is to predict how Alzheimer's disease progresses "
        "by combining clinical data, brain scans (MRI), and brain activity (EEG) over time."
    )

    # Phase 0: Audit
    doc.add_heading('2. Phase 0: Forensic Audit and Cleanup', level=1)
    doc.add_paragraph(
        "First, we audited the existing codebase to determine what was actually built versus what was merely scaffolded. "
        "We discovered that while the clinical data pipeline (using the OASIS-2 dataset) was functioning, the MRI and EEG components "
        "were incomplete stubs. We established strict rules to prevent 'hallucinating' results and ensured no fake data would be used."
    )
    p0 = doc.add_paragraph()
    p0.add_run("Key Files Created/Modified:\n").bold = True
    p0.add_run("- docs/CURRENT_STATE_AUDIT.md: Documented the true status of all components and identified data leakage risks.")

    # Phase 1: Data Rationale
    doc.add_heading('3. Phase 1: Dataset and Modality Rationale', level=1)
    doc.add_paragraph(
        "We defined the rules for how different types of data (Clinical, MRI, EEG) would be combined. We established a strict "
        "'±90 days tolerance' rule for matching MRI/EEG scans to clinical visits to prevent temporal data leakage. We also established "
        "that patients must be split into Train/Validation/Test sets completely independently, so a patient's data doesn't accidentally "
        "leak across sets."
    )
    p1 = doc.add_paragraph()
    p1.add_run("Key Files Created/Modified:\n").bold = True
    p1.add_run("- docs/DATASET_AND_MODALITY_RATIONALE.md: Established the scientific rules for multimodal alignment.")

    # Phase 2: Clinical Longitudinal Pipeline
    doc.add_heading('4. Phase 2: Clinical Longitudinal Evaluation', level=1)
    doc.add_paragraph(
        "We built and tested our first predictive models using strictly clinical and demographic data (like age, gender, education, and cognitive test scores). "
        "The models take a sequence of past patient visits and predict their Clinical Dementia Rating (CDR) at their next future visit."
    )
    
    doc.add_heading('Technologies & Models Tested:', level=2)
    doc.add_paragraph("We implemented 7 different models using Python, Scikit-Learn, and PyTorch:", style='List Bullet')
    doc.add_paragraph("Dummy Classifier (Baseline): Always predicts the most common outcome.", style='List Bullet')
    doc.add_paragraph("Last Visit Baseline: Predicts that the patient's state will remain exactly the same as their last visit.", style='List Bullet')
    doc.add_paragraph("Logistic Regression: A standard linear statistical model.", style='List Bullet')
    doc.add_paragraph("Random Forest: An ensemble of decision trees.", style='List Bullet')
    doc.add_paragraph("Clinical MLP: A standard Deep Neural Network.", style='List Bullet')
    doc.add_paragraph("GRU (Gated Recurrent Unit): A deep learning model specialized in handling sequential timeline data.", style='List Bullet')
    doc.add_paragraph("LSTM (Long Short-Term Memory): Another deep sequence model.", style='List Bullet')

    doc.add_heading('Data & Metrics:', level=2)
    doc.add_paragraph(
        "Dataset: We used the OASIS-2 longitudinal dataset containing 373 visits across 150 subjects. "
        "We tested the models based on Accuracy, Balanced Accuracy, Precision, Recall, Macro F1-Score, and Mean Absolute Error (MAE). "
        "We successfully ran an evaluation script that trained all 7 models and saved their performance metrics."
    )

    p2 = doc.add_paragraph()
    p2.add_run("Key Files Created/Modified:\n").bold = True
    p2.add_run("- src/cerebro_x/models/baselines.py: Added the traditional Machine Learning baselines.\n")
    p2.add_run("- src/cerebro_x/models/deep/temporal.py: Added the Deep Learning models (LSTM/GRU).\n")
    p2.add_run("- scripts/run_phase2_evaluation.py: Created a pipeline to automatically load data, split it safely, train all models, and evaluate them.\n")
    p2.add_run("- artifacts/experiments/PHASE2_CLINICAL/metrics.json: Saved the final performance scores for all models.")

    # Next Steps
    doc.add_heading('5. Next Steps (Phase 3)', level=1)
    doc.add_paragraph(
        "We are now moving to Phase 3: The MRI Pipeline. We will build the infrastructure to ingest actual 3D brain scans, "
        "process them (Quality Control, Normalization), and train a 3D Convolutional Neural Network (3D CNN) to extract features "
        "from the scans. If we require heavy GPU compute for this, we will move the training process to Kaggle."
    )

    # Save
    os.makedirs('artifacts', exist_ok=True)
    doc.save('artifacts/Cerebro_X_Progress_Summary.docx')
    print("Document generated at artifacts/Cerebro_X_Progress_Summary.docx")

if __name__ == "__main__":
    create_summary_doc()
