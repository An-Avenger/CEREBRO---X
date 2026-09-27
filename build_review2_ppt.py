"""
CEREBRO X – Review 2 PPT Builder
Matches the exact template of the Review 1 PPT:
  - Times New Roman font
  - VIT dark-navy title bars (from master)
  - White title text
  - VIT Chennai logo on title slide (from master)
  - Slide numbers bottom-right (from layout)
  - 21 slides total
"""

import os, sys

# ─── Colour helpers ──────────────────────────────────────────────────────────
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

SRC  = r"e:\PROJECTS\CEREBRO-X\25MCS1018 ARYAN SHARMA PROJECT WORK REVIEW 1 PPT_V2.pptx"
DEST = r"e:\PROJECTS\CEREBRO-X\25MCS1018 ARYAN SHARMA CEREBRO-X REVIEW 2 PPT.pptx"

def rgb(h):
    return RGBColor(int(h[0:2],16), int(h[2:4],16), int(h[4:6],16))

DARK_BLUE  = rgb('1F3864')   # VIT dark navy – body text
MID_BLUE   = rgb('1F497D')   # heading blue
WHITE      = rgb('FFFFFF')
GOLD       = rgb('BF8F00')   # warm gold highlights
GREEN      = rgb('375623')   # finding/detail green
ORANGE     = rgb('C55A11')   # in-progress orange
TF         = "Times New Roman"

# ─── Load source; properly delete all slides ─────────────────────────────────
prs = Presentation(SRC)
title_layout   = prs.slide_layouts[0]   # centre-title layout
content_layout = prs.slide_layouts[1]   # title + content layout

def _delete_slide(prs, idx):
    """Remove slide at *idx* – properly removes the part AND the sldId entry."""
    sldIdLst = prs.slides._sldIdLst
    sId = sldIdLst[idx]
    rId = sId.get(
        '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    )
    # Remove the relationship from the presentation part (python-pptx 1.x API)
    if rId:
        prs.part.drop_rel(rId)
    sldIdLst.remove(sId)

# Delete from last to first to keep indices stable
for i in range(len(prs.slides) - 1, -1, -1):
    _delete_slide(prs, i)

print(f"Slides remaining after deletion: {len(prs.slides)}")
assert len(prs.slides) == 0, "All slides must be deleted before building new ones"

# ─── Text-frame builder ───────────────────────────────────────────────────────
def add_run(para, text, size_pt=18, bold=False, color=DARK_BLUE, font=TF):
    run = para.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    return run

def add_para(tf, first=False):
    """Return a new paragraph (or reuse the first empty one)."""
    if first and len(tf.paragraphs) == 1 and tf.paragraphs[0].text == '':
        return tf.paragraphs[0]
    return tf.add_paragraph()

def set_spacing(para, before_pt=5, level=0):
    para.space_before = Pt(before_pt)
    para.level = level

# ─── Slide factory ────────────────────────────────────────────────────────────
def new_content_slide(title_text):
    """Add a content slide (layout 1) and return (slide, body_tf)."""
    slide = prs.slides.add_slide(content_layout)

    # --- set title ---
    title_ph = None
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 0:
            title_ph = ph
            break
    if title_ph:
        title_ph.text = title_text
        for para in title_ph.text_frame.paragraphs:
            for run in para.runs:
                run.font.name = TF
                run.font.size = Pt(36)
                run.font.bold = True
                run.font.color.rgb = WHITE

    # --- find or create body text-frame ---
    body_ph = None
    for ph in slide.placeholders:
        idx = ph.placeholder_format.idx
        if idx not in (0, 12):   # 0=title, 12=slide-number
            body_ph = ph
            break

    if body_ph is None:
        txBox = slide.shapes.add_textbox(
            Inches(0.45), Inches(1.55), Inches(13.1), Inches(5.9)
        )
        tf = txBox.text_frame
    else:
        tf = body_ph.text_frame

    tf.word_wrap = True
    return slide, tf


def heading(tf, text, size=20, color=MID_BLUE, spc=8, first=False):
    p = add_para(tf, first=first)
    set_spacing(p, before_pt=spc)
    add_run(p, text, size_pt=size, bold=True, color=color)
    return p

def body_text(tf, text, size=18, bold=False, color=DARK_BLUE, spc=5):
    p = add_para(tf)
    set_spacing(p, before_pt=spc)
    add_run(p, text, size_pt=size, bold=bold, color=color)
    return p

def multirun(tf, segments, size=18, spc=5):
    """segments = list of (text, bold, color)"""
    p = add_para(tf)
    set_spacing(p, before_pt=spc)
    for txt, bld, clr in segments:
        add_run(p, txt, size_pt=size, bold=bld, color=clr)
    return p

def bullet(tf, text, size=18, color=DARK_BLUE, spc=4):
    p = add_para(tf)
    set_spacing(p, before_pt=spc)
    add_run(p, "\u2022  " + text, size_pt=size, color=color)
    return p

def gap(tf, pt=4):
    p = add_para(tf)
    p.space_before = Pt(pt)

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 – TITLE SLIDE
# ═══════════════════════════════════════════════════════════════════════════════
slide1 = prs.slides.add_slide(title_layout)

for ph in slide1.placeholders:
    idx = ph.placeholder_format.idx
    if idx == 0:                         # centre title
        ph.text = "CEREBRO X"
        for para in ph.text_frame.paragraphs:
            for run in para.runs:
                run.font.name = TF
                run.font.size = Pt(44)
                run.font.bold = True
                run.font.color.rgb = WHITE
    elif idx == 1:                        # subtitle
        tf = ph.text_frame
        tf.clear()
        lines = [
            ("An AI-Powered Digital Brain Twin Framework for",        26, False, WHITE),
            ("Alzheimer's Disease Progression Prediction",            26, True,  WHITE),
            ("",                                                      10, False, WHITE),
            ("M.Tech Project Work  \u2013  Review 2",                 22, False, GOLD),
            ("",                                                       8, False, WHITE),
            ("Presented by:  Aryan Sharma  (25MCS1018)",              21, False, MID_BLUE),
            ("M.Tech CSE (Artificial Intelligence & Machine Learning)",21, False, MID_BLUE),
            ("",                                                       8, False, WHITE),
            ("Under the Guidance of:  Dr. V. Sakthivel",             21, False, MID_BLUE),
            ("Associate Professor",                                   21, False, MID_BLUE),
            ("School of Computer Science and Engineering",            21, False, MID_BLUE),
            ("Vellore Institute of Technology, Chennai",              21, False, MID_BLUE),
        ]
        first = True
        for txt, sz, bld, clr in lines:
            if first:
                para = tf.paragraphs[0]
                first = False
            else:
                para = tf.add_paragraph()
            para.alignment = PP_ALIGN.CENTER
            if txt:
                add_run(para, txt, size_pt=sz, bold=bld, color=clr)

print("Slide 1 (Title) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 – TABLE OF CONTENTS
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("TABLE OF CONTENTS")
heading(tf, "M.Tech Project Work \u2013 Review 2  |  Cerebro X", size=20, spc=4, first=True)
gap(tf, 6)
toc = [
    "1.    Problem Statement",
    "2.    Literature Review \u2013 Key Papers",
    "3.    Literature Review \u2013 Recent Work (2023\u20132026)",
    "4.    10-Paper Literature Table",
    "5.    Research Gap",
    "6.    Research Objectives",
    "7.    Proposed Solution",
    "8.    System Architecture",
    "9.    Dataset Description \u2013 ADNI",
    "10.  Dataset Description \u2013 OASIS-3",
    "11.  Data Preprocessing",
    "12.  Modules Identified",
    "13.  AI Role in Each Module",
    "14.  Evaluation Plan",
    "15.  Guide Approval",
    "16.  Status of Work",
    "17.  Timeline",
    "18.  Expected Contributions",
    "19.  Project Summary",
]
for t in toc:
    bullet(tf, t, size=17, spc=3)
print("Slide 2 (TOC) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 – PROBLEM STATEMENT
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("PROBLEM STATEMENT")
body_text(tf, "Alzheimer's disease is a progressive neurodegenerative disorder causing gradual cognitive decline.", size=19, spc=4)
gap(tf)
for b in [
    "Existing AI systems mainly focus on static diagnosis/classification from individual MRI scans.",
    "They fail to model patient-specific disease progression over time.",
    "MRI, clinical history and cognitive assessments are frequently analysed separately.",
    "Many deep-learning models provide predictions without sufficient interpretability.",
    "There is a need for an explainable, longitudinal and personalized AI framework.",
]:
    bullet(tf, b, size=18)
gap(tf, 8)
heading(tf, "One-Line Problem Statement:", size=19, spc=6)
body_text(tf,
    '"Existing Alzheimer\'s AI systems largely provide point-in-time predictions rather than a '
    'continuously evolving, explainable representation of an individual patient\'s disease progression."',
    size=18, bold=True, color=GOLD, spc=4)
print("Slide 3 (Problem Statement) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 – LITERATURE REVIEW: KEY PAPERS
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("LITERATURE REVIEW \u2013 KEY PAPERS")
heading(tf, "Key Literature (Selected Studies):", size=20, spc=4, first=True)
gap(tf, 6)
papers = [
    ("Lee et al., 2019",      "Multimodal RNN | ADNI",
     "Longitudinal multimodal data improved MCI\u2192AD prediction; accuracy 81%, AUC 0.86"),
    ("Li et al., 2019",       "Deep learning time-to-event | ADNI + external",
     "Predicted MCI\u2192AD progression; C-index 0.762 and 0.781"),
    ("Abrol et al., 2020",    "Deep Residual Learning | ADNI",
     "Baseline MRI-based progression prediction achieved 83% accuracy"),
    ("Basahel et al., 2021",  "Multimodal deep learning | ADNI",
     "MRI + clinical/genetic information improved AD stage classification"),
    ("DenseNet-BiLSTM, 2024", "CNN + BiLSTM | ADNI",
     "Combined spatial MRI features with longitudinal temporal information"),
    ("Aghaei et al., 2025",   "Ensemble + Generative + XAI | ADNI",
     "Predicts future AD progression from cognitively normal subjects"),
]
for author, approach, finding in papers:
    multirun(tf, [
        (f"\u25b6  {author}  \u2014  {approach}", True, MID_BLUE),
    ], size=18, spc=6)
    multirun(tf, [
        (f"       {finding}", False, GREEN),
    ], size=17, spc=2)
    gap(tf, 3)
print("Slide 4 (Lit Review Key) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 – LITERATURE REVIEW: RECENT WORK
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("LITERATURE REVIEW \u2013 RECENT WORK")
heading(tf, "Recent Research Directions (2023\u20132026):", size=20, spc=4, first=True)
gap(tf, 6)
recent = [
    ("2023", "MRI + XAI",                       "Grad-CAM can identify AD-relevant brain regions"),
    ("2024", "Longitudinal MRI + BiLSTM",        "Temporal information improves progression modelling"),
    ("2025", "MRI Radiomics + Progression",      "Longitudinal structural features useful for MCI\u2192AD prediction"),
    ("2025", "Integrated Prediction + XAI",      "Progression prediction combined with interpretable AI"),
    ("2026", "2.5D Longitudinal CNN",            "Subject-level splitting emphasized to reduce data leakage"),
    ("2026", "MRI + Clinical ML",                "Routine clinical + MRI data for 12-month cognitive decline prediction"),
    ("2026", "Multimodal MRI + Clinical XAI",    "Transformer-based fusion with Grad-CAM + SHAP"),
    ("2026", "Longitudinal Progression Modelling","Deep learning to model brain changes across disease stages"),
]
for year, direction, obs in recent:
    multirun(tf, [
        (f"[{year}]  ", True, GOLD),
        (f"{direction}:  ", True, MID_BLUE),
        (obs, False, DARK_BLUE),
    ], size=18, spc=5)
gap(tf, 6)
body_text(tf,
    "Recent studies increasingly emphasise longitudinal data, multimodal information and explainability "
    "rather than simple binary classification.",
    size=17, spc=4)
print("Slide 5 (Lit Review Recent) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 6 – 10-PAPER LITERATURE TABLE
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("10-PAPER LITERATURE TABLE")
heading(tf, "Minimum 10 Reputed Papers \u2013 Required by VIT for Review 2:", size=19, spc=4, first=True)
gap(tf, 5)
ten = [
    ("1",  "Predicting AD progression using multimodal deep learning",   "2019", "Longitudinal multimodal prediction"),
    ("2",  "Deep learning model for early prediction of AD dementia",    "2019", "MRI-based progression prediction"),
    ("3",  "Deep Residual Learning for Neuroimaging",                    "2020", "MRI-based AD progression"),
    ("4",  "Multimodal deep learning models for early AD detection",     "2021", "MRI + genetic + clinical fusion"),
    ("5",  "Longitudinal MRI + DenseNet-BiLSTM",                        "2024", "Spatial + temporal modelling"),
    ("6",  "Biomarker investigation through MRI + XAI",                 "2023", "Grad-CAM interpretability"),
    ("7",  "Integrated AD progression prediction + interpretable AI",   "2025", "Progression + XAI"),
    ("8",  "Longitudinal MRI + radiomics progression prediction",       "2025", "MRI + temporal modelling"),
    ("9",  "ML prediction of 12-month cognitive decline",               "2026", "MRI + routine clinical data"),
    ("10", "Longitudinal and explainable 2.5D framework",               "2026", "Longitudinal MRI + explainability"),
]
for no, paper, year, contrib in ten:
    multirun(tf, [
        (f"  {no:>2}.  ", True, GOLD),
        (f"{paper}  ", False, DARK_BLUE),
        (f"[{year}]  ", True, MID_BLUE),
        (contrib, False, GREEN),
    ], size=17, spc=4)
gap(tf, 5)
body_text(tf,
    "Venues: Scientific Reports | J. Neuroscience Methods | Nature Aging | npj Systems Biology | IEEE JBHI",
    size=16, spc=4)
print("Slide 6 (10-paper table) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 7 – RESEARCH GAP
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("RESEARCH GAP")
gaps = [
    ("1. Static Prediction",              "Most systems focus on diagnosis at a single time point."),
    ("2. Limited Personalization",        "Population-level models do not represent individual disease trajectory."),
    ("3. Fragmented Multimodal Analysis", "MRI, clinical and cognitive information are often processed separately."),
    ("4. Limited Longitudinal Modelling", "Many approaches do not maintain a continuously evolving patient representation."),
    ("5. Explainability Gap",             "Predictions are often not sufficiently interpretable for clinical users."),
    ("6. Generalization Gap",             "Many models are evaluated within the same dataset, not on independent cohorts."),
]
first = True
for gt, gd in gaps:
    multirun(tf, [
        (f"\u25c6  {gt}:  ", True, MID_BLUE),
        (gd, False, DARK_BLUE),
    ], size=19, spc=7 if not first else 4)
    first = False
    gap(tf, 2)
gap(tf, 8)
heading(tf, "Gap We Target:", size=19, spc=6)
body_text(tf,
    "A unified, patient-specific and explainable Digital Brain Twin integrating longitudinal MRI, "
    "clinical and cognitive information for Alzheimer's disease progression prediction remains insufficiently explored.",
    size=19, bold=True, color=GOLD, spc=4)
print("Slide 7 (Research Gap) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 8 – RESEARCH OBJECTIVES
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("RESEARCH OBJECTIVES")
heading(tf, "Primary Objective:", size=20, spc=4, first=True)
body_text(tf,
    "To develop an AI-powered Digital Brain Twin framework for personalized and explainable "
    "Alzheimer's disease progression prediction using longitudinal MRI and clinical information.",
    size=19, bold=False, color=GOLD, spc=4)
gap(tf, 6)
heading(tf, "Specific Objectives:", size=20, spc=6)
objs = [
    "Collect and preprocess longitudinal MRI, clinical and cognitive data.",
    "Extract meaningful features from MRI and patient information.",
    "Develop a multimodal AI model for disease progression prediction.",
    "Construct a patient-specific Digital Brain Twin representing disease state over time.",
    "Predict future disease progression and cognitive decline.",
    "Integrate Explainable AI using Grad-CAM and SHAP techniques.",
    "Evaluate generalisation using an independent external dataset (OASIS-3).",
    "Provide visualisation and decision-support capabilities for researchers/clinicians.",
]
for i, ob in enumerate(objs, 1):
    multirun(tf, [
        (f"  {i}.  ", True, GOLD),
        (ob, False, DARK_BLUE),
    ], size=18, spc=4)
print("Slide 8 (Objectives) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 9 – PROPOSED SOLUTION
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("PROPOSED SOLUTION")
heading(tf, "Core Idea:", size=20, spc=4, first=True)
body_text(tf,
    '"Instead of predicting only \'Does the patient have Alzheimer\'s?\', Cerebro X aims to model '
    '\'How is this patient\'s condition changing, and what is the likely future trajectory?\'"',
    size=19, bold=False, color=GOLD, spc=4)
gap(tf, 6)
heading(tf, "System Pipeline:", size=20, spc=6)
pipeline = [
    ("Longitudinal Patient Data",           "MRI + Clinical + Cognitive Information"),
    ("Data Preprocessing",                  "Quality check, normalization, alignment"),
    ("Feature Extraction",                  "CNN/Transformer for MRI + ML encoder for clinical data"),
    ("Multimodal Feature Fusion",           "Combines imaging + clinical + cognitive features"),
    ("Patient-Specific Digital Brain Twin", "Continuously evolving patient representation"),
    ("Disease Progression Prediction",      "Temporal ML / LSTM / Deep Learning"),
    ("Explainable AI",                      "Grad-CAM + SHAP \u2013 transparent predictions"),
    ("Clinical Decision Support",           "Visualisation, timeline, predictions"),
]
for i, (step, detail) in enumerate(pipeline):
    arr = "\u25b6 " if i == 0 else "\u2193 "
    multirun(tf, [
        (f"  {arr}", True, MID_BLUE),
        (f"{step}: ", True, DARK_BLUE),
        (detail, False, GREEN),
    ], size=18, spc=4)
print("Slide 9 (Proposed Solution) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 10 – SYSTEM ARCHITECTURE
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("SYSTEM ARCHITECTURE")
heading(tf, "Cerebro X \u2013 Digital Brain Twin Framework Architecture:", size=19, spc=4, first=True)
gap(tf, 5)
arch = [
    ("INPUT LAYER",         ["ADNI (Training Dataset)", "OASIS-3 (External Validation)"]),
    ("PREPROCESSING",       ["NIfTI loading, skull stripping", "Normalization, registration, augmentation"]),
    ("FEATURE EXTRACTION",  ["MRI: CNN/Transformer spatial features", "Clinical/Cognitive: ML encoder"]),
    ("MULTIMODAL FUSION",   ["Deep learning-based feature fusion", "Combines imaging + structured data"]),
    ("DIGITAL BRAIN TWIN",  ["Patient-specific longitudinal representation", "Continuously updated with new visits"]),
    ("PREDICTION",          ["Disease stage / cognitive decline prediction", "Temporal modelling across visits"]),
    ("EXPLAINABLE AI",      ["Grad-CAM: spatial MRI saliency maps", "SHAP: clinical feature importance"]),
    ("OUTPUT",              ["Disease trajectory visualisation", "Clinical decision-support dashboard"]),
]
for layer, details in arch:
    multirun(tf, [
        (f"  \u25a3  {layer}:  ", True, MID_BLUE),
        (" | ".join(details), False, DARK_BLUE),
    ], size=18, spc=5)
gap(tf, 5)
body_text(tf,
    "Note: EEG remains a future extension. Core pipeline: MRI + Clinical + Cognitive \u2192 "
    "Digital Brain Twin \u2192 Progression Prediction \u2192 XAI.",
    size=16, spc=4)
print("Slide 10 (Architecture) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 11 – DATASET: ADNI
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("DATASET DESCRIPTION \u2013 ADNI")
heading(tf, "ADNI \u2013 Alzheimer's Disease Neuroimaging Initiative", size=22, spc=4, first=True)
body_text(tf, "Role in Cerebro X:  Training / Development Dataset", size=20, bold=True, color=GOLD, spc=4)
gap(tf, 6)
heading(tf, "Data Available in ADNI:", size=19, spc=5)
for b in [
    "Structural MRI (longitudinal, multiple visits per patient)",
    "Longitudinal clinical data and medical history",
    "Cognitive assessments: MMSE, CDR, ADAS-Cog, neuropsychological scores",
    "Disease diagnosis and stage labels (CN, MCI, AD)",
    "Genetic information (APOE4 status)",
    "PET imaging data for selected participants",
]:
    bullet(tf, b, size=18, spc=4)
gap(tf, 6)
heading(tf, "Why ADNI?", size=19, spc=5)
for b in [
    "Major benchmark dataset in Alzheimer's research globally",
    "Longitudinal design enables progression modelling",
    "Multi-modal: MRI + clinical + cognitive information",
    "Allows direct comparison with published studies (Lee et al., Li et al., Abrol et al.)",
]:
    bullet(tf, b, size=18, spc=4)
print("Slide 11 (ADNI) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 12 – DATASET: OASIS-3
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("DATASET DESCRIPTION \u2013 OASIS-3")
heading(tf, "OASIS-3 \u2013 Open Access Series of Imaging Studies", size=22, spc=4, first=True)
body_text(tf, "Role in Cerebro X:  Independent External Validation Dataset", size=20, bold=True, color=GOLD, spc=4)
gap(tf, 6)
heading(tf, "Key Statistics:", size=19, spc=5)
stats = [
    ("1,098",    "total participants"),
    ("605",      "cognitively normal participants"),
    ("493",      "participants with cognitive decline"),
    ("2,000+",   "MR sessions"),
    ("1,500+",   "PET scans"),
    ("6,534",    "longitudinal clinical assessments"),
    ("3,410",    "neuropsychological assessments"),
    ("~15 years","approximate longitudinal collection period"),
]
for stat, label in stats:
    multirun(tf, [
        (f"  \u2726  {stat}  ", True, GOLD),
        (label, False, DARK_BLUE),
    ], size=19, spc=5)
gap(tf, 6)
heading(tf, "Dataset Strategy:", size=19, spc=5)
body_text(tf, "ADNI (Training)  \u2192  Cerebro X Model  \u2192  OASIS-3 (External Validation)", size=19, spc=4)
gap(tf, 4)
body_text(tf,
    "Why?  To determine whether the model learns disease-related patterns rather than "
    "dataset-specific characteristics.",
    size=19, bold=True, color=GOLD, spc=4)
print("Slide 12 (OASIS-3) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 13 – DATA PREPROCESSING
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("DATA PREPROCESSING")
heading(tf, "MRI Preprocessing Pipeline:", size=20, spc=4, first=True)
for b in [
    "NIfTI loading and format validation",
    "Image quality checking",
    "Skull stripping / brain extraction",
    "Spatial normalisation to standard atlas (MNI152)",
    "Intensity normalisation",
    "Resampling to uniform resolution",
    "Registration and alignment across visits",
    "Data augmentation (flipping, rotation, noise addition)",
]:
    bullet(tf, b, size=18, spc=4)
gap(tf, 6)
heading(tf, "Clinical Data Preprocessing:", size=20, spc=5)
for b in [
    "Missing-value handling and imputation",
    "Categorical variable encoding",
    "Feature normalisation / standardisation",
    "Temporal alignment with corresponding MRI visits",
]:
    bullet(tf, b, size=18, spc=4)
gap(tf, 5)
body_text(tf, "Cognitive Data:  MMSE | CDR | Neuropsychological Scores | ADAS-Cog", size=19, spc=4)
print("Slide 13 (Preprocessing) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 14 – MODULES IDENTIFIED
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("MODULES IDENTIFIED")
mods = [
    ("Module 1", "Data Acquisition",                      "MRI + clinical + cognitive data collection (ADNI / OASIS-3)"),
    ("Module 2", "Preprocessing",                         "Cleaning, normalisation and temporal alignment"),
    ("Module 3", "MRI Feature Extraction",                "Deep learning-based brain representation (CNN/Transformer)"),
    ("Module 4", "Clinical/Cognitive Feature Extraction", "Patient-level structured information encoding (ML/neural encoder)"),
    ("Module 5", "Multimodal Fusion",                     "Combines imaging + clinical + cognitive features (deep learning)"),
    ("Module 6", "Digital Brain Twin",                    "Maintains patient-specific longitudinal representation"),
    ("Module 7", "Progression Prediction",                "Predicts disease stage / cognitive decline over time"),
    ("Module 8", "Explainable AI",                        "Grad-CAM saliency maps / SHAP clinical feature importance"),
    ("Module 9", "Visualisation & Decision Support",      "Patient timeline, predictions and explanations"),
]
first = True
for mnum, mname, mdesc in mods:
    multirun(tf, [
        (f"  {mnum} \u2013 {mname}:  ", True, MID_BLUE),
        (mdesc, False, DARK_BLUE),
    ], size=18, spc=4 if not first else 4)
    first = False
    gap(tf, 2)
print("Slide 14 (Modules) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 15 – AI ROLE IN EACH MODULE
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("AI ROLE IN EACH MODULE")
heading(tf, "AI/ML Role Breakdown \u2013 Where is the Intelligence?", size=20, spc=4, first=True)
gap(tf, 6)
ai_roles = [
    ("MRI Feature Extraction",       "CNN / Transformer \u2013 spatial feature learning from brain MRI"),
    ("Clinical Feature Processing",  "ML / Neural encoder \u2013 structured data representation"),
    ("Multimodal Fusion",            "Deep learning fusion network"),
    ("Digital Brain Twin",           "Patient representation + longitudinal state modelling"),
    ("Progression Prediction",       "Temporal ML / LSTM / Deep Learning"),
    ("Explainable AI (XAI)",         "Grad-CAM for spatial, SHAP for clinical features"),
    ("Validation",                   "Statistical + ML evaluation metrics (Accuracy, AUC, MAE, RMSE, R\u00b2, C-index)"),
]
for mod, role in ai_roles:
    multirun(tf, [
        (f"  \u25b6  {mod}: ", True, MID_BLUE),
        (role, False, DARK_BLUE),
    ], size=19, spc=6)
print("Slide 15 (AI Role) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 16 – EVALUATION PLAN
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("EVALUATION PLAN")
heading(tf, "Classification / Stage Metrics:", size=19, spc=4, first=True)
for m in ["Accuracy", "Precision", "Recall", "F1-Score", "ROC-AUC"]:
    bullet(tf, m, size=18, spc=4)
gap(tf, 6)
heading(tf, "Progression / Regression Metrics:", size=19, spc=5)
for m in ["MAE (Mean Absolute Error)", "RMSE (Root Mean Squared Error)", "R\u00b2 Score", "Concordance Index (where applicable)"]:
    bullet(tf, m, size=18, spc=4)
gap(tf, 6)
heading(tf, "Explainability Evaluation:", size=19, spc=5)
for m in [
    "Region relevance: Grad-CAM vs. known AD-affected brain regions",
    "Clinical feature importance: SHAP consistency check",
    "Consistency with known AD-associated brain areas",
]:
    bullet(tf, m, size=18, spc=4)
gap(tf, 6)
body_text(tf,
    "Generalisation:  ADNI (Training) \u2192 OASIS-3 (External Validation) \u2013 Cross-dataset performance evaluation",
    size=19, bold=True, color=GOLD, spc=5)
print("Slide 16 (Evaluation Plan) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 17 – GUIDE APPROVAL
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("GUIDE APPROVAL")
gap(tf, 10)
heading(tf, "Project Title Approved by Project Guide", size=22, color=MID_BLUE, spc=4, first=True)
gap(tf, 10)
for line in [
    "Guide Name:        Dr. V. Sakthivel",
    "Designation:        Associate Professor",
    "Department:        School of Computer Science and Engineering",
    "Institution:           Vellore Institute of Technology, Chennai",
]:
    body_text(tf, line, size=20, spc=6)
gap(tf, 12)
body_text(tf,
    "[  GUIDE APPROVAL EMAIL SCREENSHOT / VTOP APPROVAL TO BE INSERTED HERE  ]",
    size=20, bold=True, color=GOLD, spc=8)
gap(tf, 10)
body_text(tf,
    "Note: VIT Review 2 guidelines explicitly require the Guide Approval Mail screenshot.",
    size=16, spc=5)
print("Slide 17 (Guide Approval) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 18 – STATUS OF WORK
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("STATUS OF WORK")
heading(tf, "Completed:", size=20, color=GREEN, spc=4, first=True)
done = [
    "Domain finalised: Healthcare AI / Neuroinformatics",
    "Disease finalised: Alzheimer's Disease",
    "Project title finalised: Cerebro X",
    "Problem statement identified",
    "Initial literature survey completed (10+ reputed papers)",
    "Research gaps identified",
    "Research objectives defined",
    "Proposed architecture designed",
    "Dataset strategy identified (ADNI + OASIS-3)",
    "9 Modules identified",
    "Technology stack identified",
]
for d in done:
    multirun(tf, [("[Done]  ", True, GREEN), (d, False, DARK_BLUE)], size=17, spc=3)
gap(tf, 6)
heading(tf, "In Progress:", size=20, color=ORANGE, spc=5)
prog = [
    "Detailed literature review (extended coverage)",
    "Dataset acquisition / access request",
    "Dataset preprocessing pipeline",
    "Baseline AI model design",
    "Digital Brain Twin methodology design",
]
for p in prog:
    multirun(tf, [("[WIP]  ", True, ORANGE), (p, False, DARK_BLUE)], size=17, spc=3)
print("Slide 18 (Status) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 19 – TIMELINE
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("TIMELINE")
heading(tf, "Project Timeline \u2013 VIT Review Schedule", size=20, spc=4, first=True)
gap(tf, 6)
tl = [
    ("June \u2013 July 2026",     "Domain selection, problem definition, title finalisation, initial literature survey"),
    ("July \u2013 Early Aug 2026","Literature survey completion + research gap identification"),
    ("August 2026",           "Dataset acquisition + preprocessing pipeline development"),
    ("Aug \u2013 Sep 2026",       "Baseline AI model development (MRI feature extraction + clinical encoder)"),
    ("September 2026",        "Digital Brain Twin + multimodal feature fusion"),
    ("Sep \u2013 Oct 2026",       "Progression prediction + XAI integration (Grad-CAM + SHAP)"),
    ("October 2026",          "Evaluation + results + dashboard visualisation"),
    ("October 2026",          "Research paper + thesis preparation"),
    ("28 October 2026",       "Final Review \u2013 100% implementation + results + demo"),
]
for period, work in tl:
    multirun(tf, [
        (f"  \u25b6  {period}: ", True, GOLD),
        (work, False, DARK_BLUE),
    ], size=18, spc=5)
gap(tf, 6)
body_text(tf,
    "Review 2: 12 Aug 2026  |  Review 3: 30 Sep 2026  |  Draft: 12\u201316 Oct 2026  |  Final Review: 28 Oct 2026",
    size=18, bold=True, color=MID_BLUE, spc=5)
print("Slide 19 (Timeline) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 20 – EXPECTED CONTRIBUTIONS
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("EXPECTED CONTRIBUTIONS")
heading(tf, "Expected Contributions of Cerebro X:", size=20, spc=4, first=True)
gap(tf, 6)
contribs = [
    ("1.", "Patient-specific Digital Brain Twin",       "For continuous Alzheimer's monitoring and personalised care."),
    ("2.", "Longitudinal Disease Progression Modelling","Tracking how each patient's condition changes over time."),
    ("3.", "Multimodal Integration",                   "Unified framework for MRI + clinical + cognitive information."),
    ("4.", "Explainable AI",                           "Grad-CAM + SHAP for transparent, interpretable predictions."),
    ("5.", "External Validation",                      "OASIS-3 validation to assess generalisation beyond ADNI."),
    ("6.", "Clinical Decision-Support Visualisation",  "Disease trajectory view for researchers and clinicians."),
]
for num, title, desc in contribs:
    multirun(tf, [
        (f"  {num}  {title}:  ", True, MID_BLUE),
        (desc, False, DARK_BLUE),
    ], size=19, spc=7)
print("Slide 20 (Contributions) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 21 – PROJECT SUMMARY
# ═══════════════════════════════════════════════════════════════════════════════
_, tf = new_content_slide("CEREBRO X \u2013 PROJECT SUMMARY")
heading(tf, "One-Minute Summary:", size=21, spc=4, first=True)
gap(tf, 5)
body_text(tf,
    "Cerebro X is an AI-powered Digital Brain Twin framework designed for personalised Alzheimer's "
    "disease progression prediction. The system integrates longitudinal MRI scans with clinical and "
    "cognitive information, extracts multimodal patient representations, and maintains a patient-specific "
    "Digital Brain Twin that evolves with new observations. AI models predict disease progression, while "
    "Explainable AI techniques provide interpretable evidence. ADNI is used for model development and "
    "OASIS-3 for independent external validation, evaluating both predictive performance and generalisation.",
    size=18, spc=4)
gap(tf, 8)
heading(tf, "Key Numbers to Remember:", size=19, spc=6)
knums = [
    "OASIS-3: 1,098 participants | 605 cognitively normal | 493 cognitive decline",
    "OASIS-3: 2,000+ MRI sessions | 1,500+ PET scans | 6,534 clinical assessments | 3,410 neuropsychological assessments",
    "~15 years: OASIS-3 longitudinal collection period",
    "10: Minimum papers required by VIT for Review 2",
    "Marks: 7/20 Literature + Gap | 5/20 Architecture | 3/20 Guide | 5/20 Presentation & Q&A",
]
for kn in knums:
    multirun(tf, [
        ("  \u25c6  ", True, GOLD),
        (kn, False, DARK_BLUE),
    ], size=18, spc=4)
print("Slide 21 (Summary) done")

# ═══════════════════════════════════════════════════════════════════════════════
# SAVE
# ═══════════════════════════════════════════════════════════════════════════════
prs.save(DEST)
print(f"\nSaved: {DEST}")
print(f"Total slides: {len(prs.slides)}")
