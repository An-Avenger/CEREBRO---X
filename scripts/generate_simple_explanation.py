import os
from docx import Document
from docx.shared import Pt, RGBColor

def create_simple_explanation():
    doc = Document()
    
    # Title
    title = doc.add_heading('How Cerebro-X Works', level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x2D, 0x3A, 0x40)
        
    doc.add_paragraph("A simple, plain-language guide to understanding the entire system.\n")
    
    # Section 1
    doc.add_heading('1. The Big Picture', level=1)
    doc.add_paragraph(
        "Cerebro-X is an Artificial Intelligence (AI) system designed to predict how Alzheimer's Disease will progress in a patient. "
        "Instead of just looking at a patient once, it looks at their entire medical history over multiple hospital visits to see how they are changing over time."
    )
    
    # Section 2
    doc.add_heading('2. The Data (What we feed the AI)', level=1)
    doc.add_paragraph(
        "We give the AI two main types of information from the patient's past visits:\n"
        "• Clinical Data: Things like the patient's Age, Education, and cognitive test scores (like the MMSE).\n"
        "• Brain Scans (MRI): Measurements of the brain's physical size, which shrinks as Alzheimer's gets worse."
    )
    
    # Section 3
    doc.add_heading('3. The Brain of the System (The AI Model)', level=1)
    doc.add_paragraph(
        "The core AI is called a 'Temporal GRU' (Gated Recurrent Unit). You can think of it like reading a flipbook. "
        "Instead of looking at one page (one hospital visit), it flips through all the pages (all past visits) in order. "
        "By doing this, it understands the 'story' of the patient's health and can accurately guess what the next page (their future condition) will look like."
    )
    
    # Section 4
    doc.add_heading('4. How the Pieces Fit Together', level=1)
    doc.add_paragraph(
        "The system is built in three main layers:\n"
    )
    
    doc.add_paragraph(
        "Layer 1: The Engine (AI Models)\n"
        "This is the Python code that actually learns from the data and makes predictions.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "Layer 2: The Bridge (FastAPI Backend)\n"
        "This is a server that runs in the background. It takes requests like 'Predict the future for Patient 123' and asks the AI Engine for the answer. It acts as a middleman.",
        style='List Bullet'
    )
    doc.add_paragraph(
        "Layer 3: The Dashboard (Next.js Frontend)\n"
        "This is the website/user interface. It's what the doctor or researcher actually looks at. It shows beautiful charts, patient timelines, and the Brain Health Index (BHI) without requiring them to know any programming.",
        style='List Bullet'
    )
    
    # Section 5
    doc.add_heading('5. Explainability (Why did the AI say that?)', level=1)
    doc.add_paragraph(
        "Doctors can't just blindly trust a computer. So, we added a feature called 'SHAP'. "
        "SHAP acts like a detective that explains the AI's decision. For example, it might say: 'I predicted the disease will get worse mainly because the patient's brain volume shrank significantly since their last visit, and their memory score dropped by 3 points.'"
    )
    
    # Save document
    output_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "DOCS", "How_CerebroX_Works_Simple.docx")
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Document saved successfully to: {output_path}")

if __name__ == "__main__":
    create_simple_explanation()
